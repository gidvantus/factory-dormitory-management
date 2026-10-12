"""Импорт вкладки «Места» из Excel: сборка шаблона и разбор файла.

Раскладка листа повторяет экран «Места»
(`app/frontend/src/workspace/dormitories/HostelPlaces.tsx`): у одного хостела на
каждый день приходится четыре значения — «Проживает М», «Проживает Ж»,
«Оплачено мест М», «Оплачено мест Ж». Поэтому в файле каждый хостел занимает
четыре строки (по одной на поле), а дни идут колонками:

    Хостел   | Показатель      | 01.11.2025 | 02.11.2025 | ...
    Хостел 1 | Проживает М     | 5          | 4          | ...
    Хостел 1 | Проживает Ж     | 3          | 2          | ...
    Хостел 1 | Оплачено мест М | 6          | 6          | ...
    Хостел 1 | Оплачено мест Ж | 4          | 4          | ...

Подписи полей взяты из `rows` на экране дословно: шаблон — контракт формата, и
правка подписи здесь ломает уже заполненные пользователем файлы.

Служебные строки экрана («Итого», «Свободных мест») в шаблон не попадают и при
разборе дают предупреждение: они считаются на чтении в `app/place_values.py`.
"""

from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import cast, get_args

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.place_values import place_day_values
from app.schemas.hostel import PlaceField

SHEET_TITLE = "Места"
HOSTEL_HEADER = "Хостел"
FIELD_HEADER = "Показатель"
DAY_FORMAT = "%d.%m.%Y"
MAX_VALUE = 1_000_000_000

# Порядок строк внутри блока хостела — как в `rows` на экране «Места».
FIELDS: tuple[PlaceField, ...] = get_args(PlaceField)
FIELD_LABELS: dict[PlaceField, str] = {
    "residents_m": "Проживает М",
    "residents_f": "Проживает Ж",
    "paid_m": "Оплачено мест М",
    "paid_f": "Оплачено мест Ж",
}
# Подписи расчётных строк экрана: их импортировать нельзя.
CALCULATED_LABELS = frozenset(
    {"итого", "итого м+ж", "свободных мест", "свободных мест м", "свободных мест ж"}
)
# Ячейки под таблицей, которыми шаблон подписывает месяц: не данные хостела.
SERVICE_LABELS = frozenset({"месяц", "год"})

TEMPLATE_ERROR_DETAIL = "Файл не похож на шаблон «Места» — скачайте шаблон заново"
MULTIPLE_MONTHS_DETAIL = "Шаблон содержит больше одного месяца"
VALUE_ERROR_DETAIL = "не целое неотрицательное число от 0 до 1 000 000 000"


class TemplateError(Exception):
    """Файл не является шаблоном «Места»: применять нечего, 422 с общим текстом."""


@dataclass(frozen=True)
class ImportIssue:
    """Ошибка или предупреждение разбора: строка и колонка файла, если известны."""

    message: str
    row: int | None = None
    column: str | None = None


@dataclass(frozen=True)
class ParsedRow:
    """Строка файла: хостел, поле и непустые значения по дням месяца."""

    row: int
    hostel: str
    field: PlaceField
    values: dict[date, int]


@dataclass
class ParsedImport:
    """Результат разбора файла: месяц, строки, ошибки и предупреждения."""

    month: date
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[ImportIssue] = field(default_factory=list)
    warnings: list[ImportIssue] = field(default_factory=list)

    @property
    def hostels(self) -> list[str]:
        """Названия хостелов файла в порядке появления, без повторов."""
        names: dict[str, str] = {}
        for row in self.rows:
            names.setdefault(normalize_name(row.hostel), row.hostel)
        return list(names.values())

    @property
    def cells(self) -> int:
        """Сколько непустых значений файл запишет."""
        return sum(len(row.values) for row in self.rows)


def normalize_text(value: object) -> str:
    """Текст ячейки без лишних пробелов; нестроковые значения — через `str`."""
    if value is None:
        return ""
    return " ".join(str(value).split())


def normalize_name(value: str) -> str:
    """Ключ сопоставления хостелов и полей: регистр и лишние пробелы не важны."""
    return normalize_text(value).casefold()


def month_days(month: date) -> list[date]:
    """Все дни месяца, начиная с его первого дня."""
    start = month.replace(day=1)
    next_start = date(start.year + (start.month == 12), start.month % 12 + 1, 1)
    return [start + timedelta(days=offset) for offset in range((next_start - start).days)]


def day_title(day: date) -> str:
    return day.strftime(DAY_FORMAT)


def _numeric_value(raw: object) -> int | None:
    """Целое из ячейки. `None` — пусто («значение удаляется»), `ValueError` — не число.

    Принимается и `1 234`, и `1 234,00`, и `1234.00`: шаблон заполняют в Excel,
    где разделитель разрядов и запятая как десятичный знак — обычное дело.
    """
    if raw is None:
        return None
    if isinstance(raw, bool):
        raise ValueError
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        if not raw.is_integer():
            raise ValueError
        return int(raw)
    text = normalize_text(raw).replace(" ", "")
    if not text:
        return None
    try:
        number = Decimal(text.replace(",", "."))
    except InvalidOperation as error:
        raise ValueError from error
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError
    return int(number)


def _header_day(raw: object) -> date | None:
    """Дата из заголовка колонки: `дд.мм.гггг` строкой или настоящая дата Excel."""
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = normalize_text(raw)
    if not text:
        return None
    for pattern in (DAY_FORMAT, "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _hostel_day_value(
    raw: Mapping[date, Mapping[str, int]], day: date, field: PlaceField
) -> int | None:
    """Значение дня через общий расчёт показателей «Места» — как на экране."""
    value = place_day_values(dict(raw.get(day, {}))).get(field)
    return None if value is None else int(value)


def build_template(
    month: date,
    hostels: Sequence[tuple[str, Mapping[date, Mapping[str, int]]]],
) -> bytes:
    """Собрать `.xlsx`-шаблон за месяц: строка на каждое поле каждого хостела."""
    workbook = Workbook()
    # У свежей книги активный лист есть всегда; в заглушках типов он необязателен.
    sheet = cast(Worksheet, workbook.active)
    sheet.title = SHEET_TITLE
    days = month_days(month)
    sheet.cell(1, 1, HOSTEL_HEADER)
    sheet.cell(1, 2, FIELD_HEADER)
    for index, day in enumerate(days):
        sheet.cell(1, 3 + index, day_title(day))

    row_index = 1
    for name, values in hostels:
        for place_field in FIELDS:
            row_index += 1
            sheet.cell(row_index, 1, name)
            sheet.cell(row_index, 2, FIELD_LABELS[place_field])
            for index, day in enumerate(days):
                value = _hostel_day_value(values, day, place_field)
                if value:
                    sheet.cell(row_index, 3 + index, value)

    last_column = 2 + len(days)
    _protect_table(sheet, row_index, last_column)
    # Месяц и год — отдельными ячейками под таблицей: в шапке тогда нет ни одной
    # служебной колонки, а строка 1 остаётся ровно заголовками.
    sheet.cell(row_index + 2, 1, "Месяц")
    sheet.cell(row_index + 2, 2, month.month)
    sheet.cell(row_index + 2, 3, "Год")
    sheet.cell(row_index + 2, 4, month.year)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _protect_table(sheet: Worksheet, last_row: int, last_column: int) -> None:
    """Защита листа от добавления и удаления строк и колонок.

    Заполнять шаблон можно: ячейки таблицы разблокированы, под защитой остаётся
    только шапка. Пароля нет — лист при необходимости снимают без него.
    """
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=last_column):
        for cell in row:
            cell.protection = Protection(locked=False)

    sheet.protection.sheet = True
    sheet.protection.insertRows = True
    sheet.protection.deleteRows = True
    sheet.protection.insertColumns = True
    sheet.protection.deleteColumns = True

    if last_row < 2:
        return
    numeric = DataValidation(
        type="whole",
        operator="between",
        formula1="0",
        formula2=str(MAX_VALUE),
        allow_blank=True,
    )
    sheet.add_data_validation(numeric)
    first, last = get_column_letter(3), get_column_letter(last_column)
    numeric.add(f"{first}2:{last}{last_row}")


def parse_places_workbook(content: bytes) -> ParsedImport:
    """Разобрать файл: месяц, строки, ошибки и предупреждения.

    Ошибки не обрывают разбор: собираются все сразу, чтобы пользователь поправил
    файл за один проход. `TemplateError` — файл вообще не похож на шаблон.
    """
    sheet = _open_sheet(content)
    day_columns, warnings = _day_columns(sheet)
    if not day_columns:
        raise TemplateError(TEMPLATE_ERROR_DETAIL)

    months = sorted({day.replace(day=1) for day in day_columns.values()})
    if len(months) > 1:
        raise TemplateError(MULTIPLE_MONTHS_DETAIL)
    month = months[0]

    errors: list[ImportIssue] = []
    _check_days(month, day_columns, errors)

    rows: list[ParsedRow] = []
    hostel_fields: dict[str, set[PlaceField]] = {}
    current = ""
    label_to_field = {normalize_name(label): field for field, label in FIELD_LABELS.items()}

    for index in range(2, (sheet.max_row or 0) + 1):
        name = normalize_text(sheet.cell(index, 1).value)
        label = normalize_text(sheet.cell(index, 2).value)
        if normalize_name(name) in SERVICE_LABELS:
            continue
        if not name and not label:
            continue
        if name:
            key = normalize_name(name)
            if key != normalize_name(current):
                if key in hostel_fields:
                    errors.append(
                        ImportIssue(
                            f"Хостел «{name}» указан двумя разными блоками строк",
                            index,
                            HOSTEL_HEADER,
                        )
                    )
                    current = ""
                    continue
                hostel_fields[key] = set()
                current = name
        elif not current:
            errors.append(ImportIssue("Укажите название хостела", index, HOSTEL_HEADER))
            continue
        else:
            name = current

        field = label_to_field.get(normalize_name(label))
        if field is None:
            if normalize_name(label) in CALCULATED_LABELS:
                warnings.append(
                    ImportIssue(
                        f"«{label}» считается автоматически и не импортируется",
                        index,
                        FIELD_HEADER,
                    )
                )
            elif not label:
                errors.append(ImportIssue("Укажите поле", index, FIELD_HEADER))
            else:
                errors.append(ImportIssue(f"Неизвестное поле «{label}»", index, FIELD_HEADER))
            continue

        key = normalize_name(name)
        if field in hostel_fields[key]:
            errors.append(
                ImportIssue(
                    f"Хостел «{name}», поле «{FIELD_LABELS[field]}» указано дважды",
                    index,
                    FIELD_HEADER,
                )
            )
            continue
        hostel_fields[key].add(field)
        rows.append(
            ParsedRow(
                row=index,
                hostel=name,
                field=field,
                values=_row_values(sheet, index, day_columns, errors),
            )
        )

    return ParsedImport(month=month, rows=rows, errors=errors, warnings=warnings)


def _open_sheet(content: bytes) -> Worksheet:
    """Лист «Места» из файла: любая проблема чтения — это «не шаблон»."""
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as error:  # любое повреждение файла читается как «не шаблон»
        raise TemplateError(TEMPLATE_ERROR_DETAIL) from error
    if SHEET_TITLE not in workbook.sheetnames:
        raise TemplateError(TEMPLATE_ERROR_DETAIL)
    sheet = workbook[SHEET_TITLE]
    if not isinstance(sheet, Worksheet):
        # Лист с нужным именем есть, но это не обычный лист (например, диаграмма).
        raise TemplateError(TEMPLATE_ERROR_DETAIL)
    headers = (
        normalize_name(normalize_text(sheet.cell(1, 1).value)),
        normalize_name(normalize_text(sheet.cell(1, 2).value)),
    )
    if headers != (normalize_name(HOSTEL_HEADER), normalize_name(FIELD_HEADER)):
        raise TemplateError(TEMPLATE_ERROR_DETAIL)
    return sheet


def _day_columns(sheet: Worksheet) -> tuple[dict[int, date], list[ImportIssue]]:
    """Колонки дней из шапки: номер колонки → дата, плюс предупреждения о мусоре."""
    warnings: list[ImportIssue] = []
    columns: dict[int, date] = {}
    for column in range(3, (sheet.max_column or 0) + 1):
        raw = sheet.cell(1, column).value
        label = normalize_text(raw)
        if not label:
            continue
        day = _header_day(raw)
        if day is None:
            if normalize_name(label) in CALCULATED_LABELS:
                warnings.append(
                    ImportIssue(f"Колонка «{label}» считается автоматически и пропущена", 1, label)
                )
            else:
                warnings.append(
                    ImportIssue(f"Колонка «{label}» не распознана и пропущена", 1, label)
                )
            continue
        if day in columns.values():
            warnings.append(ImportIssue(f"Колонка «{label}» указана дважды", 1, label))
            continue
        columns[column] = day
    return columns, warnings


def _check_days(month: date, day_columns: Mapping[int, date], errors: list[ImportIssue]) -> None:
    """Пропуск дня месяца — ошибка: полная замена месяца стёрла бы его данные."""
    present = set(day_columns.values())
    for day in month_days(month):
        if day not in present:
            errors.append(
                ImportIssue(
                    f"В файле нет колонки за {day_title(day)} — импорт заменил бы месяц "
                    f"и стёр данные за этот день",
                    1,
                    day_title(day),
                )
            )


def _row_values(
    sheet: Worksheet, index: int, day_columns: Mapping[int, date], errors: list[ImportIssue]
) -> dict[date, int]:
    """Непустые значения строки. Пустая ячейка и ноль одинаково означают удаление."""
    values: dict[date, int] = {}
    for column, day in day_columns.items():
        raw = sheet.cell(index, column).value
        try:
            value = _numeric_value(raw)
        except ValueError:
            errors.append(
                ImportIssue(
                    f"«{normalize_text(raw)}» — {VALUE_ERROR_DETAIL}", index, day_title(day)
                )
            )
            continue
        if value is None or value == 0:
            continue
        if value > MAX_VALUE:
            errors.append(
                ImportIssue(f"Значение {value} больше {MAX_VALUE}", index, day_title(day))
            )
            continue
        values[day] = value
    return values
