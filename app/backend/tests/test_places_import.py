"""Импорт вкладки «Места»: шаблон, разбор файла и полная замена месяца.

Файлы собираются `openpyxl` прямо в тесте: шаблоны в git не кладём, а разбор
проверяется на тех же книгах, что отдаёт ручка шаблона.
"""

from datetime import date
from io import BytesIO
from typing import cast

import pytest
from fastapi.testclient import TestClient
from httpx import Response
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.api import hostels as hostels_api
from app.places_import import (
    FIELD_HEADER,
    HOSTEL_HEADER,
    SHEET_TITLE,
    TEMPLATE_ERROR_DETAIL,
    month_days,
)
from tests.conftest import sign_in

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PLACES_EMAIL = "places-import@example.com"
NOVEMBER = date(2025, 11, 1)
FIELD_LABELS = ("Проживает М", "Проживает Ж", "Оплачено мест М", "Оплачено мест Ж")
CellValue = str | int | float | None


def create_dormitory(client: TestClient) -> int:
    response = client.post("/api/dormitories", json={"name": "Северное", "client_name": "Клиент"})
    assert response.status_code == 201
    return response.json()["id"]


def places_base(dormitory_id: int) -> str:
    return f"/api/dormitories/{dormitory_id}/hostels"


def day_headers(month: date) -> list[str]:
    return [day.strftime("%d.%m.%Y") for day in month_days(month)]


def book_bytes(
    header: list[CellValue], rows: list[list[CellValue]], sheet: str = SHEET_TITLE
) -> bytes:
    workbook = Workbook()
    # У свежей книги активный лист есть всегда; в заглушках типов он необязателен.
    page = cast(Worksheet, workbook.active)
    page.title = sheet
    for column, value in enumerate(header, start=1):
        page.cell(1, column, value)
    for index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            if value is not None:
                page.cell(index, column, value)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def places_file(
    month: date,
    entries: list[tuple[str, str, dict[str, CellValue]]],
    *,
    days: list[str] | None = None,
) -> bytes:
    """Файл «Места»: `entries` — пара «хостел, подпись поля» и значения по дням."""
    headers = day_headers(month) if days is None else days
    rows = [
        [hostel, label, *[values.get(header) for header in headers]]
        for hostel, label, values in entries
    ]
    return book_bytes([HOSTEL_HEADER, FIELD_HEADER, *headers], rows)


def sheet_of(content: bytes) -> Worksheet:
    return load_workbook(BytesIO(content))[SHEET_TITLE]


def book_content(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def download_template(client: TestClient, base: str, month: str = "2025-11-01") -> bytes:
    response = client.get(f"{base}/import/template?month={month}")
    assert response.status_code == 200, response.text
    return response.content


def upload(client: TestClient, base: str, content: bytes, filename: str = "mesta.xlsx") -> Response:
    return client.post(f"{base}/import", files={"file": (filename, content, XLSX_MIME)})


def november_values(client: TestClient, base: str) -> dict[str, dict[str, dict[str, int]]]:
    response = client.get(f"{base}?from=2025-11-01&to=2025-11-30")
    assert response.status_code == 200
    months = response.json()["months"]
    return {hostel["name"]: hostel["values"] for hostel in months[0]["hostels"]}


def test_template_carries_every_day_and_only_editable_fields(
    client: TestClient, db_session: Session
) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    hostel_id = client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"}).json()["id"]
    saved = client.put(f"{base}/{hostel_id}/cells/2025-11-05/residents_m", json={"value": 4})
    assert saved.status_code == 204

    response = client.get(f"{base}/import/template?month=2025-11-01")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MIME)
    assert response.headers["content-disposition"] == 'attachment; filename="mesta-2025-11.xlsx"'
    sheet = sheet_of(response.content)
    expected_headers = [HOSTEL_HEADER, FIELD_HEADER, *day_headers(NOVEMBER)]
    assert [cell.value for cell in sheet[1]] == expected_headers
    # Лишних колонок нет: ровно четыре поля-строки, дни и две служебные колонки.
    assert sheet.max_column == 32
    assert [sheet.cell(row, 2).value for row in range(2, 6)] == list(FIELD_LABELS)
    assert {sheet.cell(row, 1).value for row in range(2, 6)} == {"Хостел 1"}
    assert sheet.cell(2, 7).value == 4
    assert sheet.cell(2, 8).value is None
    assert sheet.cell(3, 7).value is None
    # Месяц и год — отдельными ячейками под таблицей.
    assert sheet.cell(7, 1).value == "Месяц"
    assert sheet.cell(7, 2).value == 11
    assert sheet.cell(7, 4).value == 2025


def test_template_protects_rows_and_columns_and_validates_numbers(
    client: TestClient, db_session: Session
) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"})

    sheet = sheet_of(download_template(client, base))

    assert sheet.protection.sheet is True
    assert sheet.protection.insertRows is True
    assert sheet.protection.deleteRows is True
    assert sheet.protection.insertColumns is True
    assert sheet.protection.deleteColumns is True
    validations = sheet.data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].type == "whole"
    assert (validations[0].formula1, validations[0].formula2) == ("0", "1000000000")
    assert str(validations[0].sqref) == "C2:AF5"


def test_downloaded_template_is_parsed_back_ready(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    hostel_id = client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"}).json()["id"]
    for field, value in (("residents_m", 4), ("paid_m", 6)):
        assert (
            client.put(
                f"{base}/{hostel_id}/cells/2025-11-05/{field}", json={"value": value}
            ).status_code
            == 204
        )

    response = upload(client, base, download_template(client, base))

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "ready"
    assert body["months"] == ["2025-11-01"]
    assert body["counts"] == {"hostels": 1, "cells": 2, "deleted": 0}
    assert body["errors"] == [] and body["warnings"] == []
    assert {row["field"]: row["values"] for row in body["rows"]} == {
        "residents_m": {"2025-11-05": 4},
        "paid_m": {"2025-11-05": 6},
        "residents_f": {},
        "paid_f": {},
    }


def test_values_with_spaces_and_comma_are_accepted(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"})
    workbook = load_workbook(BytesIO(download_template(client, base)))
    workbook[SHEET_TITLE].cell(2, 7, "1 234,00")
    workbook[SHEET_TITLE].cell(3, 7, "1234.00")

    response = upload(client, base, book_content(workbook))

    assert response.status_code == 200, response.text
    values = {row["field"]: row["values"] for row in response.json()["rows"]}
    assert values["residents_m"] == {"2025-11-05": 1234}
    assert values["residents_f"] == {"2025-11-05": 1234}


def test_text_in_numeric_column_fails_and_writes_nothing(
    client: TestClient, db_session: Session
) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    hostel_id = client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"}).json()["id"]
    saved = client.put(f"{base}/{hostel_id}/cells/2025-11-05/residents_m", json={"value": 4})
    assert saved.status_code == 204
    workbook = load_workbook(BytesIO(download_template(client, base)))
    workbook[SHEET_TITLE].cell(4, 7, "много")

    response = upload(client, base, book_content(workbook))

    assert response.status_code == 422, response.text
    body = response.json()
    assert body["status"] == "failed"
    assert len(body["errors"]) == 1
    error = body["errors"][0]
    assert (error["row"], error["column"]) == (4, "05.11.2025")
    assert "много" in error["message"] and "0 до 1 000 000 000" in error["message"]
    assert body["counts"]["cells"] == 1
    assert november_values(client, base)["Хостел 1"]["residents_m"] == {"2025-11-05": 4}


def test_missing_day_column_is_an_error(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"})
    workbook = load_workbook(BytesIO(download_template(client, base)))
    workbook[SHEET_TITLE].delete_cols(10)

    response = upload(client, base, book_content(workbook))

    assert response.status_code == 422, response.text
    errors = response.json()["errors"]
    assert any(error["column"] == "08.11.2025" and "стёр" in error["message"] for error in errors)


def test_two_months_and_foreign_file_are_rejected(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    two_months = book_bytes(
        [HOSTEL_HEADER, FIELD_HEADER, "01.11.2025", "01.12.2025"],
        [["Хостел 1", "Проживает М", 1, 1]],
    )

    response = upload(client, base, two_months)
    assert response.status_code == 422
    assert "больше одного месяца" in response.json()["detail"]

    response = upload(client, base, "не книга".encode())
    assert response.status_code == 422
    assert response.json()["detail"] == TEMPLATE_ERROR_DETAIL

    response = upload(client, base, book_bytes([HOSTEL_HEADER, FIELD_HEADER], [], sheet="Лист1"))
    assert response.status_code == 422
    assert response.json()["detail"] == TEMPLATE_ERROR_DETAIL

    response = upload(client, base, book_bytes([HOSTEL_HEADER], [], sheet=SHEET_TITLE))
    assert response.status_code == 422
    assert response.json()["detail"] == TEMPLATE_ERROR_DETAIL


def test_calculated_column_is_a_warning(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"})
    workbook = load_workbook(BytesIO(download_template(client, base)))
    workbook[SHEET_TITLE].cell(1, 33, "Итого М+Ж")
    workbook[SHEET_TITLE].cell(2, 33, 10)

    response = upload(client, base, book_content(workbook))

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "ready" and body["errors"] == []
    assert [warning["column"] for warning in body["warnings"]] == ["Итого М+Ж"]
    assert "автоматически" in body["warnings"][0]["message"]


def test_duplicate_hostel_block_is_an_error(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    content = places_file(
        NOVEMBER,
        [
            ("Хостел 1", "Проживает М", {"05.11.2025": 1}),
            ("Хостел 2", "Проживает М", {"05.11.2025": 2}),
            ("Хостел 1", "Проживает Ж", {"05.11.2025": 3}),
        ],
    )

    response = upload(client, base, content)

    assert response.status_code == 422
    errors = response.json()["errors"]
    assert any("двумя разными блоками" in error["message"] for error in errors)


def test_confirm_replaces_only_the_month_from_file(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    first = client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"}).json()["id"]
    client.post(base, json={"name": "Хостел 2", "month": "2025-11-01"})
    for day, field, value in (
        ("2025-11-05", "residents_m", 4),
        ("2025-11-05", "paid_m", 6),
        ("2025-12-03", "paid_m", 9),
    ):
        cell = f"{base}/{first}/cells/{day}/{field}"
        assert client.put(cell, json={"value": value}).status_code == 204

    content = places_file(
        NOVEMBER,
        [
            ("Хостел 1", "Проживает М", {"05.11.2025": 7}),
            ("Хостел 1", "Проживает Ж", {}),
            ("Хостел 1", "Оплачено мест М", {}),
            ("Хостел 1", "Оплачено мест Ж", {"05.11.2025": 2}),
            ("Хостел 2", "Проживает М", {"06.11.2025": 1}),
            ("Хостел 3", "Проживает М", {"10.11.2025": 3}),
        ],
    )
    parsed = upload(client, base, content)
    assert parsed.status_code == 200, parsed.text
    body = parsed.json()
    assert body["counts"] == {"hostels": 3, "cells": 4, "deleted": 1}

    applied = client.post(f"{base}/import/{body['import_id']}/confirm")

    assert applied.status_code == 200, applied.text
    assert applied.json() == {
        "import_id": body["import_id"],
        "month": "2025-11-01",
        "hostels": 3,
        "cells": 4,
    }
    months = client.get(f"{base}?from=2025-11-01&to=2025-12-31").json()["months"]
    november = {hostel["name"]: hostel["values"] for hostel in months[0]["hostels"]}
    assert november["Хостел 1"]["residents_m"] == {"2025-11-05": 7}
    assert november["Хостел 1"]["paid_f"] == {"2025-11-05": 2}
    assert november["Хостел 1"].get("paid_m", {}) == {}
    assert november["Хостел 2"]["residents_m"] == {"2025-11-06": 1}
    assert november["Хостел 3"]["residents_m"] == {"2025-11-10": 3}
    december = {hostel["name"]: hostel["values"] for hostel in months[1]["hostels"]}
    assert december["Хостел 1"]["paid_m"] == {"2025-12-03": 9}

    repeated = client.post(f"{base}/import/{body['import_id']}/confirm")
    assert repeated.status_code == 409
    assert "уже применён" in repeated.json()["detail"]
    assert november_values(client, base)["Хостел 1"]["residents_m"] == {"2025-11-05": 7}


def test_confirm_rejects_failed_and_unknown_jobs(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    broken = places_file(NOVEMBER, [("Хостел 1", "Неизвестное поле", {"05.11.2025": 1})])
    parsed = upload(client, base, broken)
    assert parsed.status_code == 422
    import_id = parsed.json()["import_id"]

    failed = client.post(f"{base}/import/{import_id}/confirm")
    assert failed.status_code == 409
    assert "не готов" in failed.json()["detail"]

    assert client.post(f"{base}/import/{import_id + 100}/confirm").status_code == 404


def test_upload_rejects_foreign_type_and_big_file(
    client: TestClient, db_session: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))
    client.post(base, json={"name": "Хостел 1", "month": "2025-11-01"})
    content = download_template(client, base)

    assert upload(client, base, content, filename="mesta.txt").status_code == 415

    monkeypatch.setattr(hostels_api, "MAX_IMPORT_BYTES", 10)
    too_big = upload(client, base, content)
    assert too_big.status_code == 413
    assert "20 МБ" in too_big.json()["detail"]


def test_import_requires_active_user_and_open_section(
    client: TestClient, db_session: Session
) -> None:
    assert (
        client.get("/api/dormitories/1/hostels/import/template?month=2025-11-01").status_code == 401
    )
    sign_in(client, db_session, email=PLACES_EMAIL, is_active=False)
    assert upload(client, places_base(1), b"x").status_code == 403

    sign_in(client, db_session, email="places-active@example.com")
    dormitory_id = create_dormitory(client)
    base = places_base(dormitory_id)
    parsed = upload(client, base, places_file(NOVEMBER, [("Хостел 1", "Проживает М", {})]))
    assert parsed.status_code == 200, parsed.text
    import_id = parsed.json()["import_id"]
    hidden = client.patch(
        f"/api/dormitories/{dormitory_id}/sections/places", json={"status": "hidden"}
    )
    assert hidden.status_code == 200

    assert upload(client, base, places_file(NOVEMBER, [])).status_code == 409
    assert client.post(f"{base}/import/{import_id}/confirm").status_code == 409


def test_template_month_must_be_the_first_day(client: TestClient, db_session: Session) -> None:
    sign_in(client, db_session, email=PLACES_EMAIL)
    base = places_base(create_dormitory(client))

    response = client.get(f"{base}/import/template?month=2025-11-15")

    assert response.status_code == 422
    assert response.json()["detail"] == "Укажите первый день месяца"
